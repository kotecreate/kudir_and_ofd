from pathlib import Path
from collections import Counter
from decimal import Decimal
import openpyxl
from openpyxl.utils import get_column_letter as L

for p in Path('.').glob('file-*.xlsx'):
    print('\nFILE',p.name)
    w=openpyxl.load_workbook(p,data_only=True)
    for s in w:
        print('SHEET',s.title,s.max_row,s.max_column)
        rows=list(s.values)
        if s.max_column>40:
            for i,row in enumerate(rows[:10],1):
                print(i,' | '.join(f'{L(j)}={v}' for j,v in enumerate(row,1) if v is not None))
            print('DOCUMENTS',Counter(r[13] for r in rows[8:] if len(r)>13))
            print('OPERATIONS',Counter(r[16] for r in rows[8:] if len(r)>16))
            for i,row in enumerate(rows,1):
                if i>s.max_row-5 or (len(row)>16 and 'возврат' in str(row[16]).lower()):
                    print(i,' | '.join(f'{L(j)}={v}' for j,v in enumerate(row,1) if v not in (None,'')))
        elif s.title.startswith('Раздел1') or s.title.startswith('Раздел-'):
            for i,row in enumerate(rows,1):
                if i<7 or i>s.max_row-5 or (len(row)>2 and ('налич' in str(row[2]).lower() or 'возврат' in str(row[2]).lower())):
                    print(i,' | '.join(f'{L(j)}={v}' for j,v in enumerate(row,1) if v is not None))
        else:
            print('sample',[(i,[str(v)[:100] for v in row if v is not None]) for i,row in enumerate(rows,1) if any(v is not None for v in row)][:3])
